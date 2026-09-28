from common import *
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import LineChart, BarChart, ScatterChart, Reference, Series
from openpyxl.utils import get_column_letter as CL

d = load_summary()
wb = Workbook(); F = "Arial"
H = PatternFill("solid", fgColor="154378"); HF = Font(name=F, bold=True, color="FFFFFF", size=10)
INP = Font(name=F, color="0000FF", size=10); NF = Font(name=F, size=10); BF = Font(name=F, bold=True, size=10)
thin = Side(style="thin", color="D0D7DE"); BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
def hdr(ws, r, c, t, w=None):
    x = ws.cell(r, c, t); x.fill = H; x.font = HF; x.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center"); x.border = BOX
    if w: ws.column_dimensions[CL(c)].width = w
def title(ws, t, sub=""):
    ws["A1"] = t; ws["A1"].font = Font(name=F, bold=True, size=14, color="154378")
    if sub: ws["A2"] = sub; ws["A2"].font = Font(name=F, italic=True, size=9, color="555555")
def axes(ch):
    ch.x_axis.delete = False; ch.y_axis.delete = False

# ---------------------------------------------------------------- README
ws = wb.active; ws.title = "README"
title(ws, "Results analysis: coupled water fluxes, irrigation and mulch control (Daisy)",
      "All analysis cells are live formulas on the 'Data' sheet. Blue text = model output values; black = formulas.")
info = [("Data", "All 10,190 runs (one row each). valid = 1 if not a solver failure. Factor columns parsed from the run names."),
        ("WaterBalance", "E0 rainfed balance per soil; 'other evaporation' = ET − (T + soil evap + canopy evap) by formula; stacked chart."),
        ("Trigger_E2", "Irrigation, yield and IE per soil × trigger by AVERAGEIFS; IE recomputed by formula; line chart."),
        ("TempRain_E11", "Temperature × rain matrix of irrigation (EpFactor 0.6) by AVERAGEIFS; sensitivity row by SLOPE."),
        ("Mulch_E14", "Irrigation and soil-evaporation savings vs no mulch (vff = 1) by formula; bar/line charts."),
        ("Nitrogen_E13", "N leaching fraction = SUMIFS(N leached)/SUMIFS(N fertiliser) per N rate × strategy; line chart."),
        ("Elasticity", "Log-log elasticities: LN() helper columns and SLOPE() for 5 drivers × 4 outputs."),
        ("ProductionFn", "Irrigation vs yield for all irrigated runs of E1–E4, E10 (scatter chart per soil)."),
        ("Pareto", "Counts of Pareto-optimal runs per experiment (COUNTIFS on the Data 'pareto' flag)."),
        ("", ""), ("Rules", "Solver-failure runs excluded (valid = 0). IE/IWUE shown only where irrigation ≥ 20 mm/yr.")]
for i, (a, b) in enumerate(info, 4):
    ws.cell(i, 1, a).font = BF; ws.cell(i, 2, b).font = NF
ws.column_dimensions["A"].width = 16; ws.column_dimensions["B"].width = 120

# ---------------------------------------------------------------- Data
par = set(pd.read_csv(os.path.join(OUT, "07_Derived_Data_Tables", "pareto_runs.csv")).run_id)
D = d.copy(); D["valid"] = (~D.failed).astype(int); D["pareto"] = D.run_id.isin(par).astype(int)
cols = ["experiment", "exp", "run_id", "soil", "treatment", "valid", "pareto", "P", "dT", "Ep", "trig", "z", "amt_mm", "vff", "cap", "Nrate", "wstrat", "root",
        "yield_Mg_ha", "irrigation_mm", "rain_mm", "drainage_1m_mm", "ET_mm", "transpiration_mm", "soil_evap_mm", "canopy_evap_mm",
        "water_stress_d", "N_fertiliser_kg_ha", "N_leached_kg_ha", "IE_kg_m3", "DPF", "NLF"]
ds = wb.create_sheet("Data")
for j, c in enumerate(cols, 1): hdr(ds, 1, j, c, 13 if j > 5 else 20)
for i, row in enumerate(D[cols].itertuples(index=False), 2):
    for j, v in enumerate(row, 1):
        if isinstance(v, float) and not np.isfinite(v): v = None
        elif isinstance(v, (np.floating,)): v = float(v)
        elif isinstance(v, (np.integer,)): v = int(v)
        c = ds.cell(i, j, v); c.font = INP
ds.freeze_panes = "D2"; ds.auto_filter.ref = f"A1:{CL(len(cols))}{len(D)+1}"
N = len(D) + 1; C = {c: CL(i) for i, c in enumerate(cols, 1)}
R = lambda c: f"Data!${C[c]}$2:${C[c]}${N}"

# ---------------------------------------------------------------- WaterBalance
w = wb.create_sheet("WaterBalance"); title(w, "Rainfed water balance per soil (E0), mm/yr", "Other evaporation = ET − T − soil evap − canopy evap (formula).")
for j, t in enumerate(["Soil", "Rain", "ET", "Transpiration", "Soil evap.", "Canopy evap.", "Other evap. (formula)", "Drainage 1 m", "Balance error (formula)"], 1): hdr(w, 4, j, t, 16)
for i, s in enumerate(SOILS, 5):
    w.cell(i, 1, s).font = NF
    for j, k in [(2, "rain_mm"), (3, "ET_mm"), (4, "transpiration_mm"), (5, "soil_evap_mm"), (6, "canopy_evap_mm"), (8, "drainage_1m_mm")]:
        w.cell(i, j, f'=AVERAGEIFS({R(k)},{R("exp")},"E0",{R("soil")},$A{i})').number_format = "0.0"
    w.cell(i, 7, f"=C{i}-D{i}-E{i}-F{i}").number_format = "0.0"; w.cell(i, 9, f"=B{i}-C{i}-H{i}").number_format = "0.0"
ch = BarChart(); ch.type = "col"; ch.grouping = "stacked"; ch.overlap = 100; ch.title = "Rainfed water partitioning (E0)"; ch.y_axis.title = "mm/yr"
ch.add_data(Reference(w, min_col=4, max_col=8, min_row=4, max_row=14), titles_from_data=True); ch.set_categories(Reference(w, min_col=1, min_row=5, max_row=14))
for sr, col in zip(ch.series, ["27AE60", "C0A16B", "5DADE2", "B8C4CE", "154378"]): sr.graphicalProperties.solidFill = col
ch.width, ch.height = 22, 10; axes(ch); w.add_chart(ch, "A17")

# ---------------------------------------------------------------- Trigger_E2
t2 = wb.create_sheet("Trigger_E2"); title(t2, "Trigger response (E2): irrigation (mm/yr) per soil × trigger", "AVERAGEIFS on Data; IE block = yield×1000/(irrigation×10), blank if irrigation < 20 mm.")
trigs = sorted(ok(d, "E2").trig.dropna().unique(), reverse=True)
hdr(t2, 4, 1, "Trigger (cm)", 12)
for j, s in enumerate(SOILS, 2): hdr(t2, 4, j, SOIL_LAB[s], 13); t2.cell(3, j, s).font = Font(name=F, size=8, color="888888")
for i, tg in enumerate(trigs, 5):
    t2.cell(i, 1, float(tg)).font = INP
    for j, s in enumerate(SOILS, 2):
        t2.cell(i, j, f'=IFERROR(AVERAGEIFS({R("irrigation_mm")},{R("exp")},"E2",{R("soil")},{CL(j)}$3,{R("trig")},$A{i},{R("valid")},1),"")').number_format = "0"
r0 = 5 + len(trigs) + 2; t2.cell(r0 - 1, 1, "Yield (Mg DM/ha)").font = BF
for i, tg in enumerate(trigs, r0):
    t2.cell(i, 1, f"=A{i - r0 + 5}")
    for j, s in enumerate(SOILS, 2):
        t2.cell(i, j, f'=IFERROR(AVERAGEIFS({R("yield_Mg_ha")},{R("exp")},"E2",{R("soil")},{CL(j)}$3,{R("trig")},$A{i},{R("valid")},1),"")').number_format = "0.00"
r1 = r0 + len(trigs) + 2; t2.cell(r1 - 1, 1, "IE (kg/m³) = yield×1000/(irrigation×10)").font = BF
for k in range(len(trigs)):
    i = r1 + k; t2.cell(i, 1, f"=A{5 + k}")
    for j in range(2, 12):
        L = CL(j); t2.cell(i, j, f'=IF(AND(ISNUMBER({L}{5+k}),{L}{5+k}>=20),{L}{r0+k}*1000/({L}{5+k}*10),"")').number_format = "0.00"
ch = LineChart(); ch.title = "Irrigation vs trigger, 10 soils (E2)"; ch.y_axis.title = "mm/yr"; ch.x_axis.title = "trigger (cm)"
ch.add_data(Reference(t2, min_col=2, max_col=11, min_row=4, max_row=4 + len(trigs)), titles_from_data=True); ch.set_categories(Reference(t2, min_col=1, min_row=5, max_row=4 + len(trigs)))
ch.width, ch.height = 24, 11; axes(ch); t2.add_chart(ch, "M4")

# ---------------------------------------------------------------- TempRain_E11
t11 = wb.create_sheet("TempRain_E11"); title(t11, "E11: irrigation (mm/yr) for temperature shift × rain scale (EpFactor 0.6)", "Last row: SLOPE of irrigation per °C for each rain scale.")
e11 = ok(d, "E11"); Ps = sorted(e11.P.unique()); dTs = sorted(e11.dT.unique())
hdr(t11, 4, 1, "ΔT (°C) \\ rain", 13)
for j, p in enumerate(Ps, 2): hdr(t11, 4, j, f"{p:g}", 10); t11.cell(3, j, float(p)).font = INP
for i, dt in enumerate(dTs, 5):
    t11.cell(i, 1, float(dt)).font = INP
    for j in range(2, 2 + len(Ps)):
        t11.cell(i, j, f'=AVERAGEIFS({R("irrigation_mm")},{R("exp")},"E11",{R("dT")},$A{i},{R("P")},{CL(j)}$3,{R("Ep")},0.6,{R("valid")},1)').number_format = "0"
rs = 5 + len(dTs); t11.cell(rs, 1, "mm per °C").font = BF
for j in range(2, 2 + len(Ps)): t11.cell(rs, j, f"=SLOPE({CL(j)}5:{CL(j)}{rs-1},$A$5:$A${rs-1})").number_format = "0.0"
ch = LineChart(); ch.title = "Irrigation vs temperature for each rain scale"; ch.y_axis.title = "mm/yr"; ch.x_axis.title = "ΔT (°C)"
ch.add_data(Reference(t11, min_col=2, max_col=1 + len(Ps), min_row=4, max_row=rs - 1), titles_from_data=True); ch.set_categories(Reference(t11, min_col=1, min_row=5, max_row=rs - 1))
ch.width, ch.height = 22, 10; axes(ch); t11.add_chart(ch, "A18")

# ---------------------------------------------------------------- Mulch_E14
m = wb.create_sheet("Mulch_E14"); title(m, "E14 mulch (rain ×1): means over interception capacities; savings vs no mulch (vff = 1.0) by formula")
for j, t in enumerate(["vapour flux factor", "soil evap (mm)", "irrigation (mm)", "drainage (mm)", "soil evap saved", "irrigation saved", "substitution ratio"], 1): hdr(m, 4, j, t, 15)
vffs = sorted(ok(d, "E14").vff.unique())
for i, v in enumerate(vffs, 5):
    m.cell(i, 1, float(v)).font = INP
    for j, k in [(2, "soil_evap_mm"), (3, "irrigation_mm"), (4, "drainage_1m_mm")]:
        m.cell(i, j, f'=AVERAGEIFS({R(k)},{R("exp")},"E14",{R("vff")},$A{i},{R("P")},1,{R("valid")},1)').number_format = "0.0"
last = 4 + len(vffs)
for i in range(5, last + 1):
    m.cell(i, 5, f"=B${last}-B{i}").number_format = "0.0"; m.cell(i, 6, f"=C${last}-C{i}").number_format = "0.0"
    m.cell(i, 7, f'=IF(E{i}>0,F{i}/E{i},"")').number_format = "0.00"
m.cell(last + 2, 1, "Overall substitution slope").font = BF
m.cell(last + 2, 2, f"=SLOPE(F5:F{last-1},E5:E{last-1})").number_format = "0.00"
ch = LineChart(); ch.title = "Savings vs mulch vapour flux factor"; ch.y_axis.title = "mm/yr saved"; ch.x_axis.title = "vapour flux factor"
ch.add_data(Reference(m, min_col=5, max_col=6, min_row=4, max_row=last), titles_from_data=True); ch.set_categories(Reference(m, min_col=1, min_row=5, max_row=last))
ch.width, ch.height = 20, 10; axes(ch); m.add_chart(ch, "I4")

# ---------------------------------------------------------------- Nitrogen_E13
n = wb.create_sheet("Nitrogen_E13"); title(n, "E13: N leaching fraction (%) = 100 × SUMIFS(N leached) / SUMIFS(N fertiliser), all soils")
ws_ = sorted(ok(d, "E13").wstrat.dropna().unique()); nr = sorted(ok(d, "E13").Nrate.unique())
hdr(n, 4, 1, "N rate (%)", 11)
for j, w_ in enumerate(ws_, 2): hdr(n, 4, j, w_[4:], 12); n.cell(3, j, w_).font = Font(name=F, size=8, color="888888")
for i, v in enumerate(nr, 5):
    n.cell(i, 1, float(v)).font = INP
    for j in range(2, 2 + len(ws_)):
        L = CL(j); crit = f'{R("exp")},"E13",{R("wstrat")},{L}$3,{R("Nrate")},$A{i},{R("valid")},1'
        n.cell(i, j, f"=100*SUMIFS({R('N_leached_kg_ha')},{crit})/SUMIFS({R('N_fertiliser_kg_ha')},{crit})").number_format = "0.0"
ch = LineChart(); ch.title = "N leaching fraction vs N rate per water strategy"; ch.y_axis.title = "%"; ch.x_axis.title = "N rate (% of course)"
ch.add_data(Reference(n, min_col=2, max_col=1 + len(ws_), min_row=4, max_row=4 + len(nr)), titles_from_data=True); ch.set_categories(Reference(n, min_col=1, min_row=5, max_row=4 + len(nr)))
ch.width, ch.height = 22, 10; axes(ch); n.add_chart(ch, "A18")

# ---------------------------------------------------------------- Elasticity
el = wb.create_sheet("Elasticity"); title(el, "Log-log elasticities: SLOPE(LN(output), LN(driver)) from factor-level means (formulas)")
drivers = [("E11", "P", "rain scale"), ("E11", "Ep", "EpFactor"), ("E12", "trig", "trigger"), ("E14", "vff", "mulch vff"), ("E15", "root", "rooting depth")]
outs = [("irrigation_mm", "irrigation"), ("drainage_1m_mm", "drainage"), ("soil_evap_mm", "soil evap."), ("yield_Mg_ha", "yield")]
row = 4; summ = []
for e, f, lab in drivers:
    lev = sorted(ok(d, e)[f].dropna().unique()); el.cell(row, 1, f"{e}: {lab}").font = BF
    hdr(el, row + 1, 1, "level", 10); hdr(el, row + 1, 2, "LN(|level|)", 11)
    for j, (o, ol) in enumerate(outs):
        hdr(el, row + 1, 3 + 2 * j, ol, 11); hdr(el, row + 1, 4 + 2 * j, f"LN({ol})", 11)
    for i, v in enumerate(lev):
        r = row + 2 + i; el.cell(r, 1, float(v)).font = INP; el.cell(r, 2, f"=LN(ABS(A{r}))")
        for j, (o, ol) in enumerate(outs):
            el.cell(r, 3 + 2 * j, f'=AVERAGEIFS({R(o)},{R("exp")},"{e}",{R(f)},$A{r},{R("valid")},1)').number_format = "0.00"
            el.cell(r, 4 + 2 * j, f'=IF({CL(3+2*j)}{r}>0,LN({CL(3+2*j)}{r}),"")')
    r_end = row + 1 + len(lev); summ.append((f"{e} {lab}", [f"=SLOPE({CL(4+2*j)}{row+2}:{CL(4+2*j)}{r_end},B{row+2}:B{r_end})" for j in range(len(outs))]))
    row = r_end + 2
sr = row + 1; el.cell(sr - 1, 1, "Elasticity summary").font = Font(name=F, bold=True, size=12, color="154378")
hdr(el, sr, 1, "driver", 22)
for j, (_, ol) in enumerate(outs, 2): hdr(el, sr, j, ol, 12)
for i, (lab, fs) in enumerate(summ, sr + 1):
    el.cell(i, 1, lab).font = NF
    for j, fx in enumerate(fs, 2): el.cell(i, j, fx).number_format = "+0.00;-0.00"

# ---------------------------------------------------------------- ProductionFn
pf = wb.create_sheet("ProductionFn"); title(pf, "Irrigation production function: irrigated runs E1–E4 and E10 (values), scatter per soil")
irr = pd.concat([ok(d, e) for e in ["E1", "E2", "E3", "E4", "E10"]]); irr = irr[irr.irrigation_mm >= 20]
ch = ScatterChart(); ch.title = "Yield vs irrigation, per soil"; ch.x_axis.title = "irrigation (mm/yr)"; ch.y_axis.title = "yield (Mg DM/ha)"; ch.style = 13
col = 1
for s in SOILS:
    z = irr[irr.soil == s][["irrigation_mm", "yield_Mg_ha"]].sort_values("irrigation_mm")
    if len(z) < 3: continue
    hdr(pf, 4, col, f"{SOIL_LAB[s]} I", 11); hdr(pf, 4, col + 1, "yield", 9)
    for i, (a, b) in enumerate(z.values, 5): pf.cell(i, col, round(float(a), 1)).font = INP; pf.cell(i, col + 1, round(float(b), 3)).font = INP
    sr_ = Series(Reference(pf, min_col=col + 1, min_row=5, max_row=4 + len(z)), Reference(pf, min_col=col, min_row=5, max_row=4 + len(z)), title=SOIL_LAB[s])
    sr_.marker.symbol = "circle"; sr_.marker.size = 4; sr_.graphicalProperties.line.noFill = True; ch.series.append(sr_); col += 3
ch.width, ch.height = 24, 12; axes(ch); pf.add_chart(ch, "A" + str(8 + irr.groupby("soil").size().max()) if False else "AF4")

# ---------------------------------------------------------------- Pareto
pa = wb.create_sheet("Pareto"); title(pa, "Pareto-optimal runs (max IE, min DPF, min NLF) per experiment: COUNTIFS on Data")
for j, t in enumerate(["experiment", "valid runs", "Pareto runs", "share (%)"], 1): hdr(pa, 4, j, t, 22 if j == 1 else 12)
for i, e in enumerate(sorted(d.experiment.unique(), key=lambda s: int(s.split("_")[0][1:])), 5):
    pa.cell(i, 1, e).font = NF; pa.cell(i, 2, f'=COUNTIFS({R("experiment")},A{i},{R("valid")},1)')
    pa.cell(i, 3, f'=COUNTIFS({R("experiment")},A{i},{R("pareto")},1)'); pa.cell(i, 4, f'=IF(B{i}>0,100*C{i}/B{i},0)').number_format = "0.0"
ch = BarChart(); ch.type = "bar"; ch.title = "Pareto-optimal runs per experiment"; ch.add_data(Reference(pa, min_col=3, min_row=4, max_row=21), titles_from_data=True)
ch.set_categories(Reference(pa, min_col=1, min_row=5, max_row=21)); ch.width, ch.height = 20, 11; axes(ch); pa.add_chart(ch, "F4")

p = os.path.join(OUT, "03_Excel_Analysis"); os.makedirs(p, exist_ok=True)
wb.save(os.path.join(p, "Results_Analysis.xlsx")); print("saved")
